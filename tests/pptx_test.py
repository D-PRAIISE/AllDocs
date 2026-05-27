"""PPTX round-trip test.

Builds a 2-slide PPTX with a title, body, and one picture; imports it,
exports it, re-imports, and verifies the shape structure survives.

Run from repo root:
    python tests/pptx_test.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "services" / "office-adapter" / "src"))

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import pptx  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from PIL import Image  # noqa: E402

from office_adapter.pptx_adapter import export_pptx, import_pptx  # noqa: E402


def build_fixture(path: Path, image_path: Path) -> None:
    prs = pptx.Presentation()
    blank = prs.slide_layouts[6]

    # Slide 1: title + body text box
    s1 = prs.slides.add_slide(blank)
    title = s1.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    title.text_frame.text = "AllDocs Slides"
    for run in title.text_frame.paragraphs[0].runs:
        run.font.size = Pt(36)
    body = s1.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(2))
    body.text_frame.text = "Slide one body text."

    # Slide 2: a picture
    img = Image.new("RGB", (200, 100), "red")
    img.save(image_path)
    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_picture(str(image_path), Inches(1), Inches(1), Inches(4), Inches(2))

    prs.save(str(path))


def _shape_kinds(slide: dict) -> list[str]:
    out = []
    for sh in slide.get("shapes") or []:
        layer = sh.get("layer") or {}
        out.append(layer.get("type") or "?")
    return out


def main() -> int:
    print("PPTX round-trip test\n")
    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        src = td_path / "in.pptx"
        img = td_path / "test.png"
        build_fixture(src, img)

        udm1 = import_pptx(src)
        if udm1.get("mode") != "slides":
            print(f"  [FAIL] expected mode='slides', got {udm1.get('mode')!r}")
            return 1
        slides1 = udm1["content"]["slides"]
        print(f"  imported {len(slides1)} slide(s):")
        for i, s in enumerate(slides1):
            print(f"    slide {i + 1}: shapes={_shape_kinds(s)}")

        ok = True
        if len(slides1) != 2:
            print(f"  [FAIL] expected 2 slides; got {len(slides1)}")
            ok = False

        kinds1 = _shape_kinds(slides1[0])
        if kinds1.count("text") < 2:
            print(f"  [FAIL] slide 1 should have ≥2 text shapes; got {kinds1}")
            ok = False
        kinds2 = _shape_kinds(slides1[1])
        if "image" not in kinds2:
            print(f"  [FAIL] slide 2 should contain an image; got {kinds2}")
            ok = False
        if not (udm1.get("resources") or {}).get("images"):
            print("  [FAIL] no image resource extracted")
            ok = False

        if ok:
            print("  [PASS] importer captured text + image shapes")

        # Round-trip.
        out = td_path / "out.pptx"
        export_pptx(udm1, out)
        udm2 = import_pptx(out)
        slides2 = udm2["content"]["slides"]
        if len(slides2) != len(slides1):
            print(f"  [FAIL] slide count drifted: {len(slides1)} -> {len(slides2)}")
            ok = False
        else:
            print(f"  [PASS] slide count preserved ({len(slides2)})")

        # Image should still be there on slide 2.
        kinds_after = _shape_kinds(slides2[1]) if len(slides2) > 1 else []
        if "image" not in kinds_after:
            print(f"  [FAIL] image lost on round-trip; slide 2 has {kinds_after}")
            ok = False
        else:
            print("  [PASS] image survived round-trip")

    print()
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
