"""
XLSX <-> UDM sheet-mode adapter.

Maps openpyxl's workbook model into the UDM sheet schema defined in
core/udm/src/sheet.rs:

  - workbook.sheetnames → SheetContent.sheets[].name
  - cell values         → Cell.value (tagged Number / String / Boolean / Date / Error / Empty)
  - cell formulas       → Cell.formula (the original "=..." text)
  - cell number_format  → Cell.format
  - cell comment        → Cell.comment
  - cell hyperlink.target → Cell.hyperlink
  - column widths       → Sheet.column_widths (in points; openpyxl
                          gives "character units" — we convert at ~7px/char @ 96dpi)
  - row heights         → Sheet.row_heights (already in points)
  - merged ranges       → Sheet.merges

The cells map is sparse — empty cells stay out of it. Charts, pivot
tables, conditional formats, and protection are preserved at the
zip-level via passthrough in a future stretch; v1 doesn't surface
them in UDM.
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import uuid
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

UDM_VERSION = "0.1.0"

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# openpyxl returns column widths in "character units"; 7 px/char is the
# Excel-ish default at the standard font size, then convert to points.
_CHAR_TO_PT = 7 * 72 / 96
_DEFAULT_COL_WIDTH_PT = 8.43 * _CHAR_TO_PT


def import_xlsx(path: str | Path) -> dict[str, Any]:
    """Read an XLSX file and produce a UDM v0.1.0 sheet-mode document."""
    path = Path(path)
    raw = path.read_bytes()
    sha256 = hashlib.sha256(raw).hexdigest()
    # data_only=False so we keep formulas; openpyxl gives us the
    # formula text in cell.value when this is False.
    wb = openpyxl.load_workbook(str(path), data_only=False)

    sheets: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        sheets.append(_worksheet_to_udm(ws))

    metadata: dict[str, Any] = {
        "source_format": {"mime": XLSX_MIME, "extension": "xlsx"},
        "source_hash": sha256,
    }
    if wb.properties.title:
        metadata["title"] = wb.properties.title
    if wb.properties.creator:
        metadata["authors"] = [wb.properties.creator]
    if wb.properties.subject:
        metadata["subject"] = wb.properties.subject

    return {
        "udm_version": UDM_VERSION,
        "id": str(uuid.uuid4()),
        "mode": "sheet",
        "metadata": metadata,
        "resources": {"fonts": [], "images": [], "styles": [], "colors": []},
        "content": {"sheets": sheets},
    }


def export_xlsx(udm: dict[str, Any], output_path: str | Path) -> None:
    """UDM sheet → .xlsx file."""
    if udm.get("mode") != "sheet":
        raise ValueError(f"XLSX export requires sheet UDM; got mode={udm.get('mode')!r}")
    content = udm.get("content") or {}
    sheets = content.get("sheets") or []

    wb = openpyxl.Workbook()
    # Workbook is created with a default sheet "Sheet" — remove it
    # before adding ours so we don't end up with a stray.
    default = wb.active
    if default is not None:
        wb.remove(default)

    if not sheets:
        # An entirely empty workbook would be invalid; keep one sheet.
        wb.create_sheet("Sheet1")

    for sheet_data in sheets:
        ws = wb.create_sheet(sheet_data.get("name") or "Sheet")
        for addr, cell in (sheet_data.get("cells") or {}).items():
            try:
                ws_cell = ws[addr]
            except (KeyError, ValueError):
                continue
            formula = cell.get("formula")
            value = cell.get("value") or {}
            kind = value.get("kind")
            if formula:
                # openpyxl stores formulas as strings starting with "=".
                ws_cell.value = (
                    formula if formula.startswith("=") else f"={formula}"
                )
            elif kind == "number":
                ws_cell.value = value.get("value")
            elif kind == "string":
                ws_cell.value = value.get("value")
            elif kind == "boolean":
                ws_cell.value = bool(value.get("value"))
            elif kind == "date":
                iso = value.get("value") or ""
                ws_cell.value = _parse_iso_date(iso) or iso
            elif kind == "empty":
                continue
            if cell.get("format"):
                ws_cell.number_format = cell["format"]
            if cell.get("hyperlink"):
                ws_cell.hyperlink = cell["hyperlink"]
        # Merges.
        for rng in sheet_data.get("merges") or []:
            start = rng.get("start")
            end = rng.get("end")
            if start and end:
                try:
                    ws.merge_cells(f"{start}:{end}")
                except (ValueError, AttributeError):
                    continue
        # Column widths back to character units.
        for col, length in (sheet_data.get("column_widths") or {}).items():
            pt = _length_to_pt(length)
            if pt and pt > 0:
                ws.column_dimensions[col].width = pt / _CHAR_TO_PT
        # Row heights.
        for row, length in (sheet_data.get("row_heights") or {}).items():
            pt = _length_to_pt(length)
            if pt and pt > 0:
                try:
                    ws.row_dimensions[int(row)].height = pt
                except (ValueError, TypeError):
                    continue

    wb.save(str(output_path))


# ---- import helpers ----------------------------------------------------


def _worksheet_to_udm(ws: Any) -> dict[str, Any]:
    cells: dict[str, dict[str, Any]] = {}
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            if cell.value is None and not cell.comment and not cell.hyperlink:
                continue
            addr = cell.coordinate
            entry: dict[str, Any] = {"value": _cell_value(cell)}
            # openpyxl exposes formula text in cell.value (since we
            # loaded with data_only=False) when the value starts with
            # "=". Cell.data_type == "f" is the definitive signal.
            if cell.data_type == "f" and isinstance(cell.value, str):
                entry["formula"] = cell.value
                # The "display" (cached calculated value) is lost when
                # data_only=False. We could re-open with data_only=True
                # to fetch it, but that doubles import time.
            if cell.number_format and cell.number_format != "General":
                entry["format"] = cell.number_format
            if cell.comment:
                entry["comment"] = {
                    "author": cell.comment.author or "",
                    "text": cell.comment.text or "",
                }
            if cell.hyperlink and cell.hyperlink.target:
                entry["hyperlink"] = cell.hyperlink.target
            cells[addr] = entry

    merges: list[dict[str, Any]] = []
    for merged in ws.merged_cells.ranges:
        coords = str(merged)  # "A1:B3"
        if ":" in coords:
            start, end = coords.split(":", 1)
            merges.append({"start": start, "end": end})

    col_widths: dict[str, dict[str, Any]] = {}
    for letter, dim in (ws.column_dimensions or {}).items():
        if dim.width is None or dim.width == 0:
            continue
        col_widths[letter] = _length_pt(dim.width * _CHAR_TO_PT)

    row_heights: dict[str, dict[str, Any]] = {}
    for row_idx, dim in (ws.row_dimensions or {}).items():
        if dim.height is None or dim.height == 0:
            continue
        row_heights[str(row_idx)] = _length_pt(float(dim.height))

    sheet: dict[str, Any] = {"name": ws.title, "cells": cells}
    if merges:
        sheet["merges"] = merges
    if col_widths:
        sheet["column_widths"] = col_widths
    if row_heights:
        sheet["row_heights"] = row_heights
    if ws.sheet_state == "hidden":
        sheet["hidden"] = True
    return sheet


def _cell_value(cell: Any) -> dict[str, Any]:
    """openpyxl Cell → UDM CellValue (tagged union)."""
    v = cell.value
    if v is None:
        return {"kind": "empty"}
    if cell.data_type == "f" and isinstance(v, str):
        # Formula: the value field carries the textual formula. UDM
        # stores the formula on the Cell, not as the value; emit Empty
        # here so the value field still validates.
        return {"kind": "empty"}
    if cell.data_type == "e" and isinstance(v, str):
        return {"kind": "error", "code": v}
    if isinstance(v, bool):
        return {"kind": "boolean", "value": v}
    if isinstance(v, (int, float)):
        return {"kind": "number", "value": float(v)}
    if isinstance(v, _dt.datetime):
        return {"kind": "date", "value": v.isoformat()}
    if isinstance(v, _dt.date):
        return {"kind": "date", "value": v.isoformat()}
    if isinstance(v, _dt.time):
        return {"kind": "date", "value": v.isoformat()}
    return {"kind": "string", "value": str(v)}


def _length_pt(value: float) -> dict[str, Any]:
    return {"value": float(value), "unit": "pt"}


def _length_to_pt(length: Any) -> float | None:
    if not isinstance(length, dict):
        return None
    try:
        v = float(length.get("value") or 0)
    except (TypeError, ValueError):
        return None
    unit = length.get("unit") or "pt"
    if unit == "pt":
        return v
    if unit == "px":
        return v * 72 / 96
    if unit == "in":
        return v * 72
    if unit == "mm":
        return v * 72 / 25.4
    return v


def _parse_iso_date(s: str) -> _dt.datetime | None:
    if not s:
        return None
    try:
        return _dt.datetime.fromisoformat(s)
    except ValueError:
        return None
