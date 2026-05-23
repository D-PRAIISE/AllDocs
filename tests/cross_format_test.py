"""Cross-format export test (PDF <-> DOCX).

Two round-trips:
  - PDF -> DOCX: import the whitepaper (fixed UDM), export as DOCX,
    open the DOCX and verify the original text content is present.
  - DOCX -> PDF: generate a synthetic DOCX with headings + paragraphs,
    import it (flow UDM), export as PDF, verify the text + heading
    are extractable from the output PDF.

Run from repo root:
    python tests/cross_format_test.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "services" / "pdf-adapter" / "src"))
sys.path.insert(0, str(REPO_ROOT / "services" / "office-adapter" / "src"))

import pymupdf  # noqa: E402
from docx import Document as DocxDocument  # noqa: E402

from office_adapter.exporter import export_docx  # noqa: E402
from office_adapter.importer import import_docx  # noqa: E402
from pdf_adapter.exporter import export_pdf  # noqa: E402
from pdf_adapter.importer import import_pdf  # noqa: E402

WHITEPAPER = Path.home() / "Downloads" / "1000PDF_Whitepaper.pdf"


def test_pdf_to_docx(td: Path) -> bool:
    if not WHITEPAPER.exists():
        print("  SKIP: whitepaper fixture missing")
        return True
    src = td / "src.pdf"
    src.write_bytes(WHITEPAPER.read_bytes())
    udm = import_pdf(src, source_url="/s.pdf", enable_ocr=False)

    out = td / "out.docx"
    export_docx(udm, out)
    if not out.exists() or out.stat().st_size < 1000:
        print(f"  [FAIL] PDF->DOCX wrote {out.stat().st_size if out.exists() else 0} bytes")
        return False
    doc = DocxDocument(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    ok = True
    if "1000PDF" not in text:
        print("  [FAIL] PDF->DOCX missing original text")
        ok = False
    else:
        print("  [PASS] PDF->DOCX preserved document text")
    if "INTRODUCTION" not in text and "Introduction" not in text:
        print("  [FAIL] PDF->DOCX missing INTRODUCTION heading")
        ok = False
    else:
        print("  [PASS] PDF->DOCX preserved heading text")
    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    if headings:
        print(f"  [PASS] PDF->DOCX produced {len(headings)} heading-styled paragraph(s)")
    else:
        print("  [FAIL] PDF->DOCX produced no headings")
        ok = False
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    if len(paragraphs) >= 5:
        print(f"  [PASS] PDF->DOCX produced {len(paragraphs)} paragraphs")
    else:
        print(f"  [FAIL] only {len(paragraphs)} paragraphs produced")
        ok = False
    return ok


def make_synthetic_docx(path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("Cross-Format Test", level=1)
    doc.add_paragraph(
        "AllDocs converts DOCX flow content to a paginated PDF. This is "
        "a synthetic document used to verify that round-trip."
    )
    doc.add_heading("Section Two", level=2)
    doc.add_paragraph("Some body text here for the second section.")
    doc.add_paragraph("Another paragraph with different content.")
    doc.save(str(path))


def test_docx_to_pdf(td: Path) -> bool:
    src = td / "src.docx"
    make_synthetic_docx(src)
    udm = import_docx(src)
    out = td / "out.pdf"
    export_pdf(udm, out, base_dir=td)
    if not out.exists() or out.stat().st_size < 500:
        print(f"  [FAIL] DOCX->PDF wrote {out.stat().st_size if out.exists() else 0} bytes")
        return False
    doc = pymupdf.open(str(out))
    try:
        text = "\n".join(p.get_text() for p in doc)
    finally:
        doc.close()
    ok = True
    # Hyphen may extract as U+00AD via the font's cmap; accept either
    # variant — visually the heading renders fine.
    norm = text.replace("­", "-")
    if "Cross-Format Test" in norm:
        print("  [PASS] DOCX->PDF preserved heading text")
    else:
        print("  [FAIL] DOCX->PDF missing heading 'Cross-Format Test'")
        ok = False
    if "synthetic document" in text:
        print("  [PASS] DOCX->PDF preserved paragraph body")
    else:
        print("  [FAIL] DOCX->PDF missing body text")
        ok = False
    if "Section Two" in text:
        print("  [PASS] DOCX->PDF preserved second heading")
    else:
        print("  [FAIL] DOCX->PDF missing 'Section Two'")
        ok = False
    return ok


def main() -> int:
    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        (td_path / "pd").mkdir(exist_ok=True)
        (td_path / "dp").mkdir(exist_ok=True)
        print("PDF -> DOCX:")
        a = test_pdf_to_docx(td_path / "pd")
        print()
        print("DOCX -> PDF:")
        b = test_docx_to_pdf(td_path / "dp")
        print()
        return 0 if (a and b) else 1


if __name__ == "__main__":
    print("Cross-format test\n")
    sys.exit(main())
