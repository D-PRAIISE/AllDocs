"""
UDM (flow mode) -> DOCX exporter. Inverse of importer.py.

Maps UDM blocks to python-docx elements. Inline marks become Run formatting.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph


def export_docx(udm: dict[str, Any], output_path: str | Path) -> None:
    """Render a UDM document to a DOCX file.

    Accepts both flow-mode (native) and fixed-mode (cross-format from
    PDF) UDM. For fixed UDM the text runs are grouped into paragraphs
    and headings via fixed_to_flow before rendering.
    """
    if udm.get("mode") == "fixed":
        from .fixed_to_flow import fixed_to_flow

        udm = fixed_to_flow(udm)
    elif udm.get("mode") != "flow":
        raise NotImplementedError(
            f"DOCX export supports flow and fixed UDM; got mode={udm.get('mode')!r}"
        )

    doc = Document()
    blocks = udm.get("content", {}).get("blocks", [])
    for block in blocks:
        _render_block(doc, block)

    _apply_metadata(doc, udm.get("metadata", {}))
    doc.save(str(output_path))


def export_docx_from_file(udm_path: str | Path, output_path: str | Path) -> None:
    udm_path = Path(udm_path)
    udm = json.loads(udm_path.read_text(encoding="utf-8"))
    export_docx(udm, output_path)


# ---- Blocks --------------------------------------------------------------


def _render_block(doc: Any, block: dict[str, Any]) -> None:
    btype = block.get("type")
    if btype == "heading":
        level = max(1, min(6, int(block.get("level", 1))))
        p = doc.add_heading(level=level)
        _render_inlines(p, block.get("inlines", []))
    elif btype == "paragraph":
        p = doc.add_paragraph()
        _render_inlines(p, block.get("inlines", []))
    elif btype == "list":
        _render_list(doc, block)
    elif btype == "table":
        _render_table(doc, block)
    elif btype == "code":
        # No code style in default DOCX; render as a paragraph.
        doc.add_paragraph(str(block.get("content", "")))
    elif btype == "quote":
        for inner in block.get("blocks", []):
            _render_block(doc, inner)
    elif btype == "divider":
        doc.add_paragraph("___")  # crude horizontal rule placeholder
    # image/embed: v0.2


def _render_list(doc: Any, block: dict[str, Any]) -> None:
    ordered = bool(block.get("ordered"))
    style = "List Number" if ordered else "List Bullet"
    for item in block.get("items", []):
        for ib in item.get("blocks", []):
            if ib.get("type") == "paragraph":
                p = doc.add_paragraph(style=style)
                _render_inlines(p, ib.get("inlines", []))
            else:
                _render_block(doc, ib)


def _render_table(doc: Any, block: dict[str, Any]) -> None:
    rows = block.get("rows", [])
    if not rows:
        return
    cols = max((len(r.get("cells", [])) for r in rows), default=0)
    if cols == 0:
        return
    table = doc.add_table(rows=len(rows), cols=cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row.get("cells", [])):
            if ci >= cols:
                break
            cell_obj = table.rows[ri].cells[ci]
            cell_obj.text = ""
            first = True
            for cb in cell.get("blocks", []):
                if cb.get("type") == "paragraph":
                    if first:
                        p = cell_obj.paragraphs[0]
                        first = False
                    else:
                        p = cell_obj.add_paragraph()
                    _render_inlines(p, cb.get("inlines", []))


# ---- Inlines -------------------------------------------------------------


def _render_inlines(p: DocxParagraph, inlines: list[dict[str, Any]]) -> None:
    for inline in inlines:
        itype = inline.get("type")
        if itype == "text":
            r = p.add_run(inline.get("text", ""))
            for mark in inline.get("marks", []):
                _apply_mark(r, mark)
        elif itype == "br":
            r = p.add_run()
            r.add_break()
        elif itype == "link":
            # v0.1: render the link's display text as a styled run.
            text = "".join(_inline_plain_text(i) for i in inline.get("inlines", []))
            r = p.add_run(text)
            r.font.color.rgb = RGBColor(0x00, 0x00, 0xEE)
            r.underline = True
        elif itype == "math":
            src = inline.get("source", {})
            r = p.add_run(src.get("latex") or src.get("mathml") or "")


def _apply_mark(run: Any, mark: dict[str, Any]) -> None:
    mtype = mark.get("type")
    if mtype == "bold":
        run.bold = True
    elif mtype == "italic":
        run.italic = True
    elif mtype == "underline":
        run.underline = True
    elif mtype == "strikethrough":
        run.font.strike = True
    elif mtype == "superscript":
        run.font.superscript = True
    elif mtype == "subscript":
        run.font.subscript = True
    elif mtype == "color":
        literal = (mark.get("color") or {}).get("literal") or {}
        if literal.get("kind") == "hex":
            h = str(literal.get("hex", "")).lstrip("#")
            if len(h) == 6:
                run.font.color.rgb = RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
        elif literal.get("kind") == "rgb":
            run.font.color.rgb = RGBColor(
                int(literal.get("r", 0)),
                int(literal.get("g", 0)),
                int(literal.get("b", 0)),
            )


def _inline_plain_text(inline: dict[str, Any]) -> str:
    if inline.get("type") == "text":
        return str(inline.get("text", ""))
    return ""


# ---- Metadata ------------------------------------------------------------


def _apply_metadata(doc: Any, metadata: dict[str, Any]) -> None:
    core = doc.core_properties
    if metadata.get("title"):
        core.title = str(metadata["title"])
    authors = metadata.get("authors")
    if authors:
        core.author = ", ".join(str(a) for a in authors)
    if metadata.get("subject"):
        core.subject = str(metadata["subject"])
    keywords = metadata.get("keywords")
    if keywords:
        core.keywords = ", ".join(str(k) for k in keywords)
